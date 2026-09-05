"""Declarative Profile Renderer for isolated, publication-gated Template Packs.

This module intentionally does not import the Legacy Renderer. It consumes one
already-prepared immutable snapshot and never falls back to another template.
"""

from __future__ import annotations

import hashlib
import io
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

from .report_snapshot import PreparedReportSnapshot, thaw_json
from .template_pack import TemplatePackError, inspect_template_pack

PROFILE_RENDERER_VERSION = "1.0"
_BLOCK_RENDERERS = {
    "rich_text",
    "asset_table",
    "result_table",
    "summary_table",
    "finding_table",
    "finding_sections",
    "remediation_table",
    "ioc_table",
    "field_group",
    "timeline_table",
    "mitre_table",
    "response_sections",
}
_HEADERS = {
    "hostname": "Hostname",
    "ip": "IP address",
    "os": "Operating system",
    "result": "Assessment result",
    "status": "Status",
    "type": "Type",
    "value": "Value",
    "timestamp": "Timestamp",
    "event": "Event",
    "evidence": "Evidence",
    "technique_id": "Technique ID",
    "technique": "Technique",
    "tactic": "Tactic",
}


class ProfileRendererError(ValueError):
    """Raised when a published profile cannot be rendered exactly as declared."""

    def __init__(
        self,
        message: str,
        *,
        semantic: str = "",
        anchor_kind: str = "",
        anchor_value: str = "",
    ):
        super().__init__(message)
        self.semantic = semantic
        self.anchor_kind = anchor_kind
        self.anchor_value = anchor_value

    def public(self) -> dict[str, str]:
        return {
            "code": "PROFILE_RENDER_FAILED",
            "message": str(self),
            "semantic": self.semantic,
            "anchorKind": self.anchor_kind,
            "anchorValue": self.anchor_value,
        }


class ProfileRenderCancelled(ProfileRendererError):
    """Raised when a caller requests cooperative cancellation."""


@dataclass(frozen=True)
class ProfileRenderResult:
    document: DocxDocument
    manifest: dict[str, Any]


@dataclass(frozen=True)
class _AnchorTarget:
    kind: str
    value: str
    element: Any


def render_profile_document(
    pack_bytes: bytes,
    prepared: PreparedReportSnapshot,
    *,
    check_cancelled: Callable[[], Any] | None = None,
    on_progress: Callable[[int, str], Any] | None = None,
) -> ProfileRenderResult:
    """Render a published Template Pack from the accepted normalized snapshot."""

    return _render_profile_document(
        pack_bytes,
        prepared,
        require_publishable=True,
        check_cancelled=check_cancelled,
        on_progress=on_progress,
    )


def render_profile_validation_candidate(
    pack_bytes: bytes,
    prepared: PreparedReportSnapshot,
    *,
    check_cancelled: Callable[[], Any] | None = None,
    on_progress: Callable[[int, str], Any] | None = None,
) -> ProfileRenderResult:
    """Render one non-publishable candidate for the isolated validation runner."""

    return _render_profile_document(
        pack_bytes,
        prepared,
        require_publishable=False,
        check_cancelled=check_cancelled,
        on_progress=on_progress,
    )


def _render_profile_document(
    pack_bytes: bytes,
    prepared: PreparedReportSnapshot,
    *,
    require_publishable: bool,
    check_cancelled: Callable[[], Any] | None,
    on_progress: Callable[[int, str], Any] | None,
) -> ProfileRenderResult:

    try:
        inspection = inspect_template_pack(pack_bytes, require_publishable=require_publishable)
    except TemplatePackError as exc:
        raise ProfileRendererError(f"Template Pack failed inspection: {exc}") from exc
    if not require_publishable:
        if inspection.publishable or inspection.profile.get("status") != "mapping_complete":
            raise ProfileRendererError(
                "Validation rendering requires a non-publishable mapping_complete candidate."
            )
        if not inspection.validation.mapping_complete:
            raise ProfileRendererError("Validation candidate does not have 100% semantic mapping.")
    accepted = prepared.accepted
    if inspection.profile["reportType"] != accepted.report_type:
        raise ProfileRendererError(
            "Template Pack report type does not match the accepted report snapshot."
        )
    if accepted.template_hash != inspection.template_sha256:
        raise ProfileRendererError(
            "Template Pack template does not match the template pinned by the accepted snapshot."
        )
    _check_cancelled(check_cancelled)
    try:
        document = Document(io.BytesIO(inspection.template_bytes))
    except (ValueError, OSError) as exc:
        raise ProfileRendererError("Template Pack DOCX could not be opened.") from exc
    payload = thaw_json(prepared.payload)
    if not isinstance(payload, dict):
        raise ProfileRendererError("Prepared report payload must be an object.")
    sources = build_profile_semantic_sources(prepared, payload)
    rendered: list[str] = []
    row_counts: dict[str, int] = {}
    slots = inspection.profile["slots"]
    total = max(1, len(slots))
    for index, slot in enumerate(slots, start=1):
        _check_cancelled(check_cancelled)
        semantic = slot["semantic"]
        anchor = slot["anchor"]
        try:
            target = _locate_anchor(document, anchor["kind"], anchor["value"])
            source_value = sources.get(slot["source"])
            row_counts[semantic] = _render_slot(
                document,
                target,
                slot,
                source_value,
                check_cancelled=check_cancelled,
            )
        except ProfileRendererError as exc:
            if not exc.semantic:
                exc.semantic = semantic
                exc.anchor_kind = anchor["kind"]
                exc.anchor_value = anchor["value"]
            raise
        rendered.append(semantic)
        if on_progress is not None:
            on_progress(round(index / total * 100), semantic)
    manifest = {
        "renderer": "profile",
        "rendererVersion": PROFILE_RENDERER_VERSION,
        "packId": inspection.pack_id,
        "packVersion": inspection.version,
        "packSha256": hashlib.sha256(pack_bytes).hexdigest(),
        "templateSha256": inspection.template_sha256,
        "requestSignature": accepted.request_signature,
        "contentSignature": prepared.content_signature,
        "reportType": accepted.report_type,
        "renderedSemantics": rendered,
        "rowCounts": row_counts,
        "legacyRendererUsed": False,
    }
    setattr(document, "_reporter_profile_manifest", manifest)
    return ProfileRenderResult(document=document, manifest=manifest)


def save_profile_report_atomic(
    result: ProfileRenderResult,
    target: Path,
    *,
    check_cancelled: Callable[[], Any] | None = None,
) -> Path:
    """Save via an owned temporary file and remove it on cancellation or failure."""

    destination = target.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}-",
        suffix=".tmp.docx",
        dir=destination.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        _check_cancelled(check_cancelled)
        result.document.save(temporary)
        _check_cancelled(check_cancelled)
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return destination


def _render_slot(
    document: DocxDocument,
    target: _AnchorTarget,
    slot: dict[str, Any],
    value: Any,
    *,
    check_cancelled: Callable[[], Any] | None,
) -> int:
    renderer = slot["renderer"]
    if renderer == "text":
        _replace_anchor_text(target, _plain_text(value))
        return 1
    if renderer not in _BLOCK_RENDERERS:
        raise ProfileRendererError(f"Unsupported profile renderer: {renderer}.")
    elements: list[Any]
    row_count = 0
    if renderer == "rich_text":
        paragraphs = _paragraph_values(value)
        elements = [_new_paragraph(document, item) for item in paragraphs]
        row_count = len(paragraphs)
    elif renderer in {
        "asset_table",
        "result_table",
        "remediation_table",
        "ioc_table",
        "timeline_table",
        "mitre_table",
    }:
        rows = value if isinstance(value, list) else []
        headers, table_rows = _mapped_rows(slot.get("fields", []), rows)
        elements = [
            _new_table(
                document,
                headers,
                table_rows,
                semantic=slot["semantic"],
                check_cancelled=check_cancelled,
            )
        ]
        row_count = len(table_rows)
    elif renderer == "summary_table":
        summary = value if isinstance(value, dict) else {}
        table_rows = [[_humanize(key), _plain_text(item)] for key, item in summary.items()]
        elements = [
            _new_table(document, ["Metric", "Value"], table_rows, semantic=slot["semantic"])
        ]
        row_count = len(table_rows)
    elif renderer == "finding_table":
        rows = value if isinstance(value, list) else []
        fields = ("hostname", "finding", "severity", "evidence", "remediation")
        elements = [
            _new_table(
                document,
                [_HEADERS.get(field, _humanize(field)) for field in fields],
                [[_plain_text(item.get(field)) for field in fields] for item in rows],
                semantic=slot["semantic"],
                check_cancelled=check_cancelled,
            )
        ]
        row_count = len(rows)
    elif renderer == "field_group":
        group = value if isinstance(value, dict) else {}
        table_rows = [[_humanize(key), _plain_text(item)] for key, item in group.items()]
        elements = [_new_table(document, ["Field", "Value"], table_rows, semantic=slot["semantic"])]
        row_count = len(table_rows)
    elif renderer == "finding_sections":
        rows = value if isinstance(value, list) else []
        elements = []
        for item in rows:
            _check_cancelled(check_cancelled)
            title = _plain_text(item.get("hostname") or item.get("finding") or "Finding")
            elements.append(_new_paragraph(document, title, bold=True))
            details = (
                ("Finding", item.get("finding")),
                ("Severity", item.get("severity")),
                ("Evidence", item.get("evidence")),
                ("Analysis", item.get("details")),
                ("Remediation", item.get("remediation")),
            )
            for label, detail in details:
                value_text = _plain_text(detail).strip()
                if value_text:
                    elements.append(_new_paragraph(document, f"{label}: {value_text}"))
        row_count = len(rows)
    else:  # response_sections
        group = value if isinstance(value, dict) else {}
        elements = []
        for key, items in group.items():
            elements.append(_new_paragraph(document, _humanize(key), bold=True))
            for item in _paragraph_values(items):
                _check_cancelled(check_cancelled)
                elements.append(_new_paragraph(document, item))
                row_count += 1
    if not elements:
        elements = [_new_paragraph(document, "No data available.")]
    _replace_anchor_blocks(target, elements)
    return row_count


def build_profile_semantic_sources(
    prepared: PreparedReportSnapshot,
    payload: dict[str, Any],
) -> dict[str, Any]:
    accepted = prepared.accepted
    servers = [item for item in payload.get("servers", []) if isinstance(item, dict)]
    clients = [item for item in payload.get("clients", []) if isinstance(item, dict)]
    metadata = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
    incident = metadata.get("incident") if isinstance(metadata.get("incident"), dict) else metadata
    findings = _flatten_findings((*servers, *clients))
    affected = [
        {
            "hostname": item.get("hostname", ""),
            "ip": item.get("ip", ""),
            "status": _assessment_label(item),
        }
        for item in (*servers, *clients)
        if _is_affected(item)
    ]
    ioc_metadata = {
        **metadata,
        "iocs": incident.get("iocs", metadata.get("iocs", [])),
    }
    iocs = _collect_iocs(ioc_metadata, (*servers, *clients))
    recommendations = metadata.get("recommendations") or _recommendations(findings)
    overview = metadata.get("overview") or (
        f"Assessed {len(servers)} server(s) and {len(clients)} client(s); "
        f"recorded {len(findings)} evidence-backed finding(s)."
    )
    return {
        "metadata.title": accepted.title,
        "analysis.overview": overview,
        "assets.servers": servers,
        "assets.clients": clients,
        "analysis.server_results": servers,
        "analysis.client_results": clients,
        "analysis.affected_assets": affected,
        "analysis.findings": findings,
        "analysis.iocs": iocs,
        "analysis.recommendations": recommendations,
        "analysis.summary": {
            "servers": len(servers),
            "clients": len(clients),
            "assets": len(servers) + len(clients),
            "findings": len(findings),
            "affectedAssets": len(affected),
        },
        "incident.info": {
            key: incident.get(key, "")
            for key in ("incident_id", "severity", "status", "detected_at")
        },
        "incident.executive_summary": incident.get("executive_summary", ""),
        "incident.affected_assets": incident.get("affected_assets", affected),
        "incident.timeline": incident.get("timeline", []),
        "incident.mitre": incident.get("mitre", incident.get("mitreMappings", [])),
        "incident.response": {
            "containment": incident.get("containment", incident.get("containmentActions", [])),
            "eradication": incident.get("eradication", incident.get("eradicationActions", [])),
            "recovery": incident.get("recovery", incident.get("recoveryActions", [])),
        },
        "incident.lessons_learned": incident.get("lessons_learned", ""),
    }


def _flatten_findings(assets: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for asset in assets:
        for finding in asset.get("findings", []):
            if not isinstance(finding, dict) or not finding.get("evidence"):
                continue
            evidence = finding.get("evidence", [])
            result.append(
                {
                    "hostname": asset.get("hostname", ""),
                    "finding": finding.get("name", finding.get("ruleId", "Finding")),
                    "severity": finding.get("severity", ""),
                    "evidence": "; ".join(
                        _plain_text(item.get("value", item))
                        if isinstance(item, dict)
                        else str(item)
                        for item in evidence
                    ),
                    "remediation": finding.get("remediation", ""),
                    "details": finding.get("description", ""),
                }
            )
    return result


def _collect_iocs(
    metadata: dict[str, Any], assets: Iterable[dict[str, Any]]
) -> list[dict[str, Any]]:
    raw: list[Any] = []
    if isinstance(metadata.get("iocs"), list):
        raw.extend(metadata["iocs"])
    for asset in assets:
        if isinstance(asset.get("iocs"), list):
            raw.extend(asset["iocs"])
    seen: set[tuple[str, str]] = set()
    result: list[dict[str, Any]] = []
    for item in raw:
        if isinstance(item, dict):
            ioc_type = _plain_text(item.get("type"))
            value = _plain_text(item.get("value"))
        else:
            ioc_type, value = "value", _plain_text(item)
        key = (ioc_type.casefold(), value.casefold())
        if value and key not in seen:
            seen.add(key)
            result.append({"type": ioc_type, "value": value})
    return result


def _mapped_rows(
    fields: list[dict[str, str]],
    rows: list[Any],
) -> tuple[list[str], list[list[str]]]:
    ordered = sorted(fields, key=lambda item: int(item["target"].split(":", 1)[1]))
    sources = [item["source"] for item in ordered]
    headers = [_HEADERS.get(source, _humanize(source)) for source in sources]
    values = [
        [_plain_text(row.get(source)) for source in sources]
        for row in rows
        if isinstance(row, dict)
    ]
    return headers, values


def _locate_anchor(document: DocxDocument, kind: str, value: str) -> _AnchorTarget:
    body = document.element.body
    matches: list[Any] = []
    if kind == "token":
        for paragraph in body.iter(qn("w:p")):
            if value in "".join(node.text or "" for node in paragraph.iter(qn("w:t"))):
                matches.append(paragraph)
    elif kind == "bookmark":
        matches = [
            item for item in body.iter(qn("w:bookmarkStart")) if item.get(qn("w:name")) == value
        ]
    elif kind == "content_control":
        for control in body.iter(qn("w:sdt")):
            tags = list(control.iter(qn("w:tag")))
            if any(tag.get(qn("w:val")) == value for tag in tags):
                matches.append(control)
    else:
        raise ProfileRendererError(f"Unsupported anchor kind: {kind}.")
    if len(matches) != 1:
        raise ProfileRendererError(
            f"Anchor must occur exactly once in the document body; found {len(matches)}."
        )
    return _AnchorTarget(kind=kind, value=value, element=matches[0])


def _replace_anchor_text(target: _AnchorTarget, text: str) -> None:
    if target.kind == "token":
        _replace_token_nodes(target.element, target.value, text)
    elif target.kind == "bookmark":
        _replace_bookmark_content(target.element, text)
    else:
        content = _sdt_content(target.element)
        for child in list(content):
            content.remove(child)
        content.append(_paragraph_xml(text))


def _replace_anchor_blocks(target: _AnchorTarget, elements: list[Any]) -> None:
    if target.kind == "content_control":
        content = _sdt_content(target.element)
        for child in list(content):
            content.remove(child)
        for element in elements:
            content.append(element)
        return
    if target.kind == "token":
        _replace_token_nodes(target.element, target.value, "")
        anchor = target.element
    else:
        _replace_bookmark_content(target.element, "")
        anchor = _ancestor(target.element, qn("w:p"))
    if anchor is None:
        raise ProfileRendererError("Block anchor is not inside a paragraph.")
    for element in elements:
        anchor.addnext(element)
        anchor = element


def _replace_token_nodes(paragraph: Any, token: str, replacement: str) -> None:
    nodes = list(paragraph.iter(qn("w:t")))
    combined = "".join(node.text or "" for node in nodes)
    start = combined.find(token)
    if start < 0 or combined.find(token, start + len(token)) >= 0:
        raise ProfileRendererError("Token anchor is missing or ambiguous during replacement.")
    end = start + len(token)
    offset = 0
    inserted = False
    for node in nodes:
        original = node.text or ""
        node_start, node_end = offset, offset + len(original)
        offset = node_end
        if node_end <= start or node_start >= end:
            continue
        prefix = original[: max(0, start - node_start)] if node_start <= start else ""
        suffix = original[max(0, end - node_start) :] if node_end >= end else ""
        node.text = prefix + (replacement if not inserted else "") + suffix
        inserted = True


def _replace_bookmark_content(start: Any, text: str) -> None:
    bookmark_id = start.get(qn("w:id"))
    parent = start.getparent()
    end = next(
        (
            item
            for item in parent
            if item.tag == qn("w:bookmarkEnd") and item.get(qn("w:id")) == bookmark_id
        ),
        None,
    )
    if end is None:
        raise ProfileRendererError("Bookmark end marker was not found in the same paragraph.")
    start_index = parent.index(start)
    end_index = parent.index(end)
    for item in list(parent)[start_index + 1 : end_index]:
        parent.remove(item)
    if text:
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
        parent.insert(parent.index(end), run)


def _sdt_content(control: Any) -> Any:
    contents = list(control.iter(qn("w:sdtContent")))
    if len(contents) != 1:
        raise ProfileRendererError("Content control must contain one sdtContent element.")
    return contents[0]


def _new_paragraph(document: DocxDocument, text: str, *, bold: bool = False) -> Any:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    return paragraph._p


def _new_table(
    document: DocxDocument,
    headers: list[str],
    rows: list[list[str]],
    *,
    semantic: str = "",
    check_cancelled: Callable[[], Any] | None = None,
) -> Any:
    effective_headers = headers or ["Value"]
    table: Table = document.add_table(rows=1, cols=len(effective_headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    if semantic:
        properties = table._tbl.tblPr
        caption = OxmlElement("w:tblCaption")
        caption.set(qn("w:val"), f"ReporterPro:{semantic}")
        properties.append(caption)
    for index, header in enumerate(effective_headers):
        table.rows[0].cells[index].text = header
    for row_index, values in enumerate(rows, start=1):
        if row_index % 100 == 0:
            _check_cancelled(check_cancelled)
        cells = table.add_row().cells
        for index in range(len(effective_headers)):
            cells[index].text = values[index] if index < len(values) else ""
    return table._tbl


def _paragraph_xml(text: str) -> Any:
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph.append(run)
    return paragraph


def _ancestor(element: Any, tag: str) -> Any | None:
    current = element
    while current is not None:
        if current.tag == tag:
            return current
        current = current.getparent()
    return None


def _paragraph_values(value: Any) -> list[str]:
    if isinstance(value, str):
        return [item.strip() for item in value.splitlines() if item.strip()] or [""]
    if isinstance(value, list):
        return [_plain_text(item) for item in value]
    if isinstance(value, dict):
        return [f"{_humanize(key)}: {_plain_text(item)}" for key, item in value.items()]
    return [_plain_text(value)] if value is not None else []


def _plain_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (str, int, float)):
        return str(value)
    if isinstance(value, list):
        return "; ".join(_plain_text(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{_humanize(key)}: {_plain_text(item)}" for key, item in value.items())
    return str(value)


def _humanize(value: str) -> str:
    return str(value).replace("_", " ").replace(".", " ").strip().title()


def _assessment_label(asset: dict[str, Any]) -> str:
    assessment = asset.get("assessment")
    if isinstance(assessment, dict) and assessment.get("label"):
        return _plain_text(assessment["label"])
    return _plain_text(asset.get("result"))


def _is_affected(asset: dict[str, Any]) -> bool:
    assessment = asset.get("assessment")
    if isinstance(assessment, dict) and assessment.get("classification") == "anomaly":
        return True
    return any(
        isinstance(item, dict) and item.get("classification") == "anomaly" and item.get("evidence")
        for item in asset.get("findings", [])
    )


def _recommendations(findings: list[dict[str, Any]]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for finding in findings:
        remediation = _plain_text(finding.get("remediation")).strip()
        if remediation and remediation.casefold() not in seen:
            seen.add(remediation.casefold())
            result.append(remediation)
    return result


def _check_cancelled(callback: Callable[[], Any] | None) -> None:
    if callback is None:
        return
    try:
        result = callback()
    except ProfileRenderCancelled:
        raise
    except Exception as exc:
        raise ProfileRenderCancelled("Profile rendering was cancelled.") from exc
    if result is True:
        raise ProfileRenderCancelled("Profile rendering was cancelled.")
