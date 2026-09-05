"""Read-only structural analysis for templates entering Template Studio.

The analyzer reports facts and conservative anchor suggestions.  It never
creates an approved mapping and has no dependency on the Legacy Renderer.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from typing import Any
from xml.etree import ElementTree

from docx import Document

from .template_analyzer import TOKEN_PATTERN, validate_docx_bytes
from .template_profiles import PROFILE_SCHEMA_VERSION, REPORT_REQUIREMENTS

MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_EXPANDED_SIZE = 100 * 1024 * 1024
MAX_DOCUMENT_XML_SIZE = 20 * 1024 * 1024
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NAMESPACE}}}"


def analyze_profile_template(data: bytes, report_type: str) -> dict[str, Any]:
    """Return immutable template facts and an unapproved mapping checklist."""

    validate_docx_bytes(data)
    if report_type not in REPORT_REQUIREMENTS:
        raise ValueError("Unsupported report type for Template Profile analysis.")
    document_xml, package_facts = _read_document_package_safely(data)
    upper_xml = document_xml.upper()
    if b"<!DOCTYPE" in upper_xml or b"<!ENTITY" in upper_xml:
        raise ValueError("Word document XML declarations are not allowed.")
    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:
        raise ValueError("Word document XML is invalid.") from exc

    content_controls = _content_controls(root)
    bookmarks = _bookmarks(root)
    tokens = _tokens(root)
    token_counts = _token_counts(root)
    anchors = [
        *({"kind": "content_control", **item} for item in content_controls),
        *({"kind": "bookmark", **item} for item in bookmarks),
        *({"kind": "token", "value": value} for value in tokens),
    ]
    document = Document(io.BytesIO(data))
    headings = []
    for index, paragraph in enumerate(document.paragraphs):
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading") and paragraph.text.strip():
            match = re.search(r"(\d+)$", style)
            headings.append(
                {
                    "index": index,
                    "level": int(match.group(1)) if match else 0,
                    "text": paragraph.text.strip(),
                }
            )
    tables = []
    for index, table in enumerate(document.tables):
        headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
        tables.append({"index": index, "headers": headers, "rowCount": len(table.rows)})

    sections = [_section_facts(section, index) for index, section in enumerate(document.sections)]
    headers, footers = _header_footer_facts(document)
    relationship_counts: dict[str, int] = {}
    external_relationships = 0
    for relationship in document.part.rels.values():
        kind = relationship.reltype.rsplit("/", 1)[-1]
        relationship_counts[kind] = relationship_counts.get(kind, 0) + 1
        if relationship.is_external:
            external_relationships += 1
    page_breaks = sum(
        1 for node in root.iter(f"{W}br") if node.get(f"{W}type", "") == "page"
    ) + sum(1 for _ in root.iter(f"{W}pageBreakBefore"))
    toc_fields = sum(
        1 for node in root.iter(f"{W}instrText") if "TOC" in (node.text or "").upper().split()
    )

    checklist = []
    for semantic in REPORT_REQUIREMENTS[report_type]:
        suggestions = _anchor_suggestions(semantic, anchors)
        checklist.append(
            {
                "semantic": semantic,
                "status": "unmapped",
                "required": True,
                "suggestedAnchors": suggestions,
            }
        )
    return {
        "schemaVersion": PROFILE_SCHEMA_VERSION,
        "reportType": report_type,
        "templateSha256": hashlib.sha256(data).hexdigest(),
        "facts": {
            "contentControls": content_controls,
            "bookmarks": bookmarks,
            "tokens": tokens,
            "tokenOccurrences": token_counts,
            "anchorConflicts": _anchor_conflicts(content_controls, bookmarks, token_counts),
            "headings": headings,
            "tables": tables,
            "sections": sections,
            "headers": headers,
            "footers": footers,
            "pageBreakCount": page_breaks,
            "tocFieldCount": toc_fields,
            "numbering": package_facts["numbering"],
            "relationships": {
                "counts": relationship_counts,
                "externalCount": external_relationships,
            },
        },
        "mapping": {
            "coveragePercent": 0.0,
            "approvedCount": 0,
            "requiredCount": len(checklist),
            "checklist": checklist,
        },
        "profileSeed": {
            "schemaVersion": PROFILE_SCHEMA_VERSION,
            "reportType": report_type,
            "status": "analyzed",
            "templateSha256": hashlib.sha256(data).hexdigest(),
            "slots": [],
        },
    }


def _read_document_package_safely(data: bytes) -> tuple[bytes, dict[str, Any]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES:
                raise ValueError("DOCX contains too many package entries.")
            if sum(info.file_size for info in infos) > MAX_DOCX_EXPANDED_SIZE:
                raise ValueError("DOCX expands beyond the analysis safety limit.")
            try:
                info = archive.getinfo("word/document.xml")
            except KeyError as exc:
                raise ValueError("DOCX is missing word/document.xml.") from exc
            if info.file_size > MAX_DOCUMENT_XML_SIZE:
                raise ValueError("Word document XML exceeds the analysis safety limit.")
            document_xml = archive.read(info)
            numbering = {"abstractDefinitionCount": 0, "instanceCount": 0}
            if "word/numbering.xml" in archive.namelist():
                numbering_info = archive.getinfo("word/numbering.xml")
                if numbering_info.file_size > MAX_DOCUMENT_XML_SIZE:
                    raise ValueError("Word numbering XML exceeds the analysis safety limit.")
                numbering_xml = archive.read(numbering_info)
                upper_numbering = numbering_xml.upper()
                if b"<!DOCTYPE" in upper_numbering or b"<!ENTITY" in upper_numbering:
                    raise ValueError("Word numbering XML declarations are not allowed.")
                try:
                    numbering_root = ElementTree.fromstring(numbering_xml)
                except ElementTree.ParseError as exc:
                    raise ValueError("Word numbering XML is invalid.") from exc
                numbering = {
                    "abstractDefinitionCount": sum(
                        1 for _ in numbering_root.iter(f"{W}abstractNum")
                    ),
                    "instanceCount": sum(1 for _ in numbering_root.iter(f"{W}num")),
                }
            return document_xml, {"numbering": numbering}
    except zipfile.BadZipFile as exc:
        raise ValueError("DOCX is not a valid Office ZIP package.") from exc


def _content_controls(root: ElementTree.Element) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    by_value: dict[str, dict[str, Any]] = {}
    for control in root.iter(f"{W}sdt"):
        properties = control.find(f"{W}sdtPr")
        if properties is None:
            continue
        tag_node = properties.find(f"{W}tag")
        alias_node = properties.find(f"{W}alias")
        tag = tag_node.get(f"{W}val", "").strip() if tag_node is not None else ""
        alias = alias_node.get(f"{W}val", "").strip() if alias_node is not None else ""
        value = tag or alias
        if value:
            if value in by_value:
                by_value[value]["occurrences"] += 1
                continue
            item: dict[str, Any] = {"value": value, "occurrences": 1}
            if tag:
                item["tag"] = tag
            if alias:
                item["alias"] = alias
            results.append(item)
            by_value[value] = item
    return results


def _bookmarks(root: ElementTree.Element) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    by_value: dict[str, dict[str, Any]] = {}
    for node in root.iter(f"{W}bookmarkStart"):
        name = node.get(f"{W}name", "").strip()
        if name and not name.startswith("_"):
            if name in by_value:
                by_value[name]["occurrences"] += 1
                continue
            item = {"value": name, "occurrences": 1}
            results.append(item)
            by_value[name] = item
    return results


def _tokens(root: ElementTree.Element) -> list[str]:
    found: set[str] = set()
    for paragraph in root.iter(f"{W}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
        found.update(TOKEN_PATTERN.findall(text))
    return sorted(found)


def _token_counts(root: ElementTree.Element) -> dict[str, int]:
    counts: dict[str, int] = {}
    for paragraph in root.iter(f"{W}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
        for token in TOKEN_PATTERN.findall(text):
            counts[token] = counts.get(token, 0) + 1
    return dict(sorted(counts.items()))


def _anchor_conflicts(
    content_controls: list[dict[str, Any]],
    bookmarks: list[dict[str, Any]],
    token_counts: dict[str, int],
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for kind, items in (("content_control", content_controls), ("bookmark", bookmarks)):
        for item in items:
            grouped.setdefault(_normalize(item["value"]), []).append(
                {
                    "kind": kind,
                    "value": item["value"],
                    "occurrences": item["occurrences"],
                }
            )
    for token, occurrences in token_counts.items():
        grouped.setdefault(_normalize(token), []).append(
            {"kind": "token", "value": token, "occurrences": occurrences}
        )
    return [
        {"normalizedValue": key, "anchors": items}
        for key, items in sorted(grouped.items())
        if len(items) > 1 or any(item["occurrences"] > 1 for item in items)
    ]


def _section_facts(section: Any, index: int) -> dict[str, Any]:
    start_type = getattr(section.start_type, "name", str(section.start_type))
    orientation = getattr(section.orientation, "name", str(section.orientation))
    return {
        "index": index,
        "startType": start_type,
        "orientation": orientation,
        "pageWidth": int(section.page_width or 0),
        "pageHeight": int(section.page_height or 0),
        "differentFirstPage": bool(section.different_first_page_header_footer),
    }


def _header_footer_facts(document: Any) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    headers: list[dict[str, Any]] = []
    footers: list[dict[str, Any]] = []
    seen_headers: set[str] = set()
    seen_footers: set[str] = set()
    for index, section in enumerate(document.sections):
        for target, story, seen, kind in (
            (headers, section.header, seen_headers, "default"),
            (headers, section.first_page_header, seen_headers, "first_page"),
            (headers, section.even_page_header, seen_headers, "even_page"),
            (footers, section.footer, seen_footers, "default"),
            (footers, section.first_page_footer, seen_footers, "first_page"),
            (footers, section.even_page_footer, seen_footers, "even_page"),
        ):
            if not story._has_definition:
                continue
            part = story._definition
            part_name = str(part.partname)
            if part_name in seen:
                continue
            seen.add(part_name)
            root = part.element
            target.append(
                {
                    "sectionIndex": index,
                    "part": part_name,
                    "kind": kind,
                    "linkedToPrevious": False,
                    "paragraphCount": sum(1 for _ in root.iter(f"{W}p")),
                    "tableCount": sum(1 for _ in root.iter(f"{W}tbl")),
                    "hasText": any((node.text or "").strip() for node in root.iter(f"{W}t")),
                }
            )
    return headers, footers


def _anchor_suggestions(semantic: str, anchors: list[dict[str, str]]) -> list[dict[str, str]]:
    normalized_semantic = _normalize(semantic)
    suggestions = []
    for anchor in anchors:
        normalized_anchor = _normalize(anchor["value"])
        if normalized_anchor in {normalized_semantic, f"reporter{normalized_semantic}"}:
            suggestions.append(
                {
                    "kind": anchor["kind"],
                    "value": anchor["value"],
                    "confidence": "exact_name",
                    "approval": "required",
                }
            )
    return suggestions


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())
