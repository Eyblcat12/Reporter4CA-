"""Backend-owned fixture, integrity and structural validation for Template Packs.

The runner operates on a non-publishable candidate pack.  It does not expose an
API, select a catalog version, or call the Legacy Renderer.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Callable
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn

from .profile_renderer import (
    ProfileRenderCancelled,
    build_profile_semantic_sources,
    render_profile_validation_candidate,
)
from .report_snapshot import PreparedReportSnapshot, thaw_json
from .template_pack import inspect_template_pack
from .template_pack_catalog import (
    TemplatePackValidationEvidence,
    build_template_pack_candidate,
)

VALIDATOR_VERSION = "template-pack-validator/1.0"
_TOKEN_PATTERN = re.compile(r"\{\{[A-Z][A-Z0-9_]{1,126}\}\}")
_TABLE_RENDERERS = frozenset(
    {
        "asset_table",
        "result_table",
        "summary_table",
        "finding_table",
        "remediation_table",
        "ioc_table",
        "field_group",
        "timeline_table",
        "mitre_table",
    }
)


class TemplatePackValidationError(ValueError):
    """Raised when a validation run or approval violates the trust contract."""


@dataclass(frozen=True)
class ValidationIssue:
    code: str
    message: str
    semantic: str = ""
    path: str = ""
    expected: Any = None
    actual: Any = None

    def public(self) -> dict[str, Any]:
        result = {
            "code": self.code,
            "message": self.message,
            "semantic": self.semantic,
            "path": self.path,
        }
        if self.expected is not None:
            result["expected"] = self.expected
        if self.actual is not None:
            result["actual"] = self.actual
        return result


@dataclass(frozen=True)
class TemplatePackValidationRun:
    run_id: str
    fixture_id: str
    report_type: str
    candidate_pack_sha256: str
    template_sha256: str
    request_signature: str
    content_signature: str
    artifact_sha256: str
    structural_sha256: str
    baseline_sha256: str
    fixture_passed: bool
    integrity_passed: bool
    structural_passed: bool
    issues: tuple[ValidationIssue, ...]
    manifest: dict[str, Any]
    structural_snapshot: dict[str, Any]
    artifact_bytes: bytes = field(repr=False, compare=False)

    @property
    def ready_for_visual_review(self) -> bool:
        return self.fixture_passed and self.integrity_passed and self.structural_passed

    def public(self) -> dict[str, Any]:
        return {
            "runId": self.run_id,
            "fixtureId": self.fixture_id,
            "reportType": self.report_type,
            "candidatePackSha256": self.candidate_pack_sha256,
            "templateSha256": self.template_sha256,
            "requestSignature": self.request_signature,
            "contentSignature": self.content_signature,
            "artifactSha256": self.artifact_sha256,
            "structuralSha256": self.structural_sha256,
            "baselineSha256": self.baseline_sha256,
            "fixturePassed": self.fixture_passed,
            "integrityPassed": self.integrity_passed,
            "structuralPassed": self.structural_passed,
            "readyForVisualReview": self.ready_for_visual_review,
            "issues": [issue.public() for issue in self.issues],
            "manifest": self.manifest,
        }


def run_template_pack_validation(
    workspace: dict[str, Any],
    template_bytes: bytes,
    prepared: PreparedReportSnapshot,
    *,
    fixture_id: str,
    structural_baseline: dict[str, Any] | None,
    check_cancelled: Callable[[], Any] | None = None,
    on_progress: Callable[[int, str], Any] | None = None,
) -> TemplatePackValidationRun:
    """Render one candidate and verify fixture, content and structural coverage."""

    normalized_fixture_id = fixture_id.strip()
    if not normalized_fixture_id or len(normalized_fixture_id) > 128:
        raise TemplatePackValidationError("Validation requires a fixture id (1-128 chars).")
    candidate = build_template_pack_candidate(workspace, template_bytes)
    inspection = inspect_template_pack(candidate, require_publishable=False)
    progress = _progress_adapter(on_progress)
    progress(5, "candidate.inspected")
    rendered = render_profile_validation_candidate(
        candidate,
        prepared,
        check_cancelled=check_cancelled,
        on_progress=lambda percent, semantic: progress(5 + round(percent * 0.65), semantic),
    )
    _check_cancelled(check_cancelled)
    output = io.BytesIO()
    rendered.document.save(output)
    artifact = output.getvalue()
    artifact_sha256 = _sha256(artifact)
    progress(75, "artifact.saved")

    payload = thaw_json(prepared.payload)
    if not isinstance(payload, dict):
        raise TemplatePackValidationError("Prepared fixture payload must be an object.")
    sources = build_profile_semantic_sources(prepared, payload)
    fixture_issues = _verify_fixture_counts(inspection.profile["slots"], sources, rendered.manifest)
    integrity_issues = _verify_rendered_content(
        artifact,
        inspection.profile["slots"],
        sources,
        rendered.manifest,
    )
    progress(85, "integrity.checked")

    snapshot = structural_snapshot(artifact, rendered.manifest)
    structural_sha256 = _canonical_sha256(snapshot)
    structural_issues: list[ValidationIssue] = []
    baseline_sha256 = ""
    if structural_baseline is None:
        structural_issues.append(
            ValidationIssue(
                "structural.baseline_missing",
                "A reviewed structural baseline is required before evidence can be issued.",
                path="structuralBaseline",
            )
        )
    else:
        baseline_sha256 = _canonical_sha256(structural_baseline)
        structural_issues.extend(structural_diff(structural_baseline, snapshot))
    progress(95, "structure.compared")

    issues = tuple((*fixture_issues, *integrity_issues, *structural_issues))
    run_identity = {
        "validator": VALIDATOR_VERSION,
        "fixtureId": normalized_fixture_id,
        "candidatePackSha256": _sha256(candidate),
        "artifactSha256": artifact_sha256,
        "structuralSha256": structural_sha256,
        "baselineSha256": baseline_sha256,
        "requestSignature": prepared.accepted.request_signature,
        "contentSignature": prepared.content_signature,
    }
    progress(100, "validation.completed")
    return TemplatePackValidationRun(
        run_id=_canonical_sha256(run_identity),
        fixture_id=normalized_fixture_id,
        report_type=prepared.accepted.report_type,
        candidate_pack_sha256=_sha256(candidate),
        template_sha256=inspection.template_sha256,
        request_signature=prepared.accepted.request_signature,
        content_signature=prepared.content_signature,
        artifact_sha256=artifact_sha256,
        structural_sha256=structural_sha256,
        baseline_sha256=baseline_sha256,
        fixture_passed=not fixture_issues,
        integrity_passed=not integrity_issues,
        structural_passed=not structural_issues,
        issues=issues,
        manifest=rendered.manifest,
        structural_snapshot=snapshot,
        artifact_bytes=artifact,
    )


def approve_template_pack_validation(
    run: TemplatePackValidationRun,
    *,
    reviewer: str,
    artifact_sha256: str,
    reviewed_at: str = "",
) -> TemplatePackValidationEvidence:
    """Issue builder evidence only for the exact artifact a human reviewed."""

    identity = reviewer.strip()
    if not identity or len(identity) > 128:
        raise TemplatePackValidationError("Visual approval requires a reviewer identity.")
    if artifact_sha256.strip().lower() != run.artifact_sha256:
        raise TemplatePackValidationError("Reviewed artifact checksum does not match this run.")
    if not run.ready_for_visual_review:
        raise TemplatePackValidationError(
            "Validation run has unresolved fixture/integrity/diff issues."
        )
    timestamp = reviewed_at.strip() or _timestamp()
    return TemplatePackValidationEvidence(
        fixture_id=run.fixture_id,
        fixture_passed=True,
        integrity_passed=True,
        visual_approved=True,
        validator_version=VALIDATOR_VERSION,
        validated_at=timestamp,
        validation_run_id=run.run_id,
        artifact_sha256=run.artifact_sha256,
        structural_sha256=run.structural_sha256,
        reviewed_by=identity,
        reviewed_at=timestamp,
    )


def structural_snapshot(docx_bytes: bytes, manifest: dict[str, Any]) -> dict[str, Any]:
    """Return a deterministic, content-aware DOCX structure snapshot."""

    document = Document(io.BytesIO(docx_bytes))
    paragraphs = [
        {
            "style": getattr(paragraph.style, "name", "") or "",
            "text": _text(paragraph.text),
            "numbered": bool(paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None),
        }
        for paragraph in document.paragraphs
        if _text(paragraph.text)
    ]
    tables: list[dict[str, Any]] = []
    for index, table in enumerate(document.tables):
        rows = [[_text(cell.text) for cell in row.cells] for row in table.rows]
        caption = table._tbl.tblPr.find(qn("w:tblCaption"))
        caption_value = caption.get(qn("w:val"), "") if caption is not None else ""
        semantic = caption_value.removeprefix("ReporterPro:") if caption_value else ""
        tables.append(
            {
                "index": index,
                "semantic": semantic,
                "rows": len(rows),
                "columns": max((len(row) for row in rows), default=0),
                "contentSha256": _canonical_sha256(rows),
                "formatSha256": _sha256(table._tbl.tblPr.xml.encode("utf-8")),
            }
        )
    sections = [
        {
            "orientation": str(section.orientation),
            "width": int(section.page_width or 0),
            "height": int(section.page_height or 0),
            "topMargin": int(section.top_margin or 0),
            "bottomMargin": int(section.bottom_margin or 0),
        }
        for section in document.sections
    ]
    body_text = "\n".join(item["text"] for item in paragraphs)
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        media = [
            {"name": name.rsplit("/", 1)[-1], "sha256": _sha256(archive.read(name))}
            for name in sorted(archive.namelist())
            if name.startswith("word/media/") and not name.endswith("/")
        ]
        relationships = _relationships(archive)
    return {
        "snapshotSchemaVersion": 1,
        "reportType": manifest.get("reportType", ""),
        "renderedSemantics": list(manifest.get("renderedSemantics", [])),
        "semanticRowCounts": dict(manifest.get("rowCounts", {})),
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": sections,
        "relationships": relationships,
        "media": media,
        "remainingTokens": sorted(set(_TOKEN_PATTERN.findall(body_text))),
    }


def structural_diff(
    expected: Any,
    actual: Any,
    path: str = "root",
    *,
    limit: int = 200,
) -> list[ValidationIssue]:
    """Return bounded, readable structural differences with semantic context."""

    differences: list[ValidationIssue] = []

    def compare(left: Any, right: Any, current: str, semantic: str = "") -> None:
        if len(differences) >= limit:
            return
        if isinstance(left, dict) and isinstance(right, dict):
            keys = sorted(set(left) | set(right))
            next_semantic = semantic or str(left.get("semantic") or right.get("semantic") or "")
            for key in keys:
                child = f"{current}.{key}"
                if key not in left:
                    differences.append(
                        ValidationIssue(
                            "structural.added",
                            "Structure was added.",
                            next_semantic,
                            child,
                            None,
                            right[key],
                        )
                    )
                elif key not in right:
                    differences.append(
                        ValidationIssue(
                            "structural.removed",
                            "Structure was removed.",
                            next_semantic,
                            child,
                            left[key],
                            None,
                        )
                    )
                else:
                    compare(left[key], right[key], child, next_semantic)
            return
        if isinstance(left, list) and isinstance(right, list):
            for index in range(max(len(left), len(right))):
                child = f"{current}[{index}]"
                if index >= len(left):
                    differences.append(
                        ValidationIssue(
                            "structural.added",
                            "List item was added.",
                            semantic,
                            child,
                            None,
                            right[index],
                        )
                    )
                elif index >= len(right):
                    differences.append(
                        ValidationIssue(
                            "structural.removed",
                            "List item was removed.",
                            semantic,
                            child,
                            left[index],
                            None,
                        )
                    )
                else:
                    compare(left[index], right[index], child, semantic)
            return
        if left != right:
            differences.append(
                ValidationIssue(
                    "structural.changed", "Structure changed.", semantic, current, left, right
                )
            )

    compare(expected, actual, path)
    if len(differences) >= limit:
        differences.append(
            ValidationIssue(
                "structural.diff_truncated",
                f"Structural diff was truncated after {limit} items.",
                path=path,
            )
        )
    return differences


def _verify_fixture_counts(
    slots: list[dict[str, Any]],
    sources: dict[str, Any],
    manifest: dict[str, Any],
) -> list[ValidationIssue]:
    issues: list[ValidationIssue] = []
    actual_counts = manifest.get("rowCounts", {})
    expected_semantics = [slot["semantic"] for slot in slots]
    if manifest.get("renderedSemantics") != expected_semantics:
        issues.append(
            ValidationIssue(
                "fixture.semantic_coverage",
                "Rendered semantic order does not match the candidate profile.",
                path="manifest.renderedSemantics",
                expected=expected_semantics,
                actual=manifest.get("renderedSemantics"),
            )
        )
    for slot in slots:
        semantic = slot["semantic"]
        expected = _source_row_count(slot["renderer"], sources.get(slot["source"]))
        actual = actual_counts.get(semantic)
        if actual != expected:
            issues.append(
                ValidationIssue(
                    "fixture.row_count",
                    "Rendered row count does not match the fixture source.",
                    semantic,
                    f"manifest.rowCounts.{semantic}",
                    expected,
                    actual,
                )
            )
    return issues


def _verify_rendered_content(
    artifact: bytes,
    slots: list[dict[str, Any]],
    sources: dict[str, Any],
    manifest: dict[str, Any],
) -> list[ValidationIssue]:
    document = Document(io.BytesIO(artifact))
    document_text = "\n".join(node.text or "" for node in document.element.body.iter(qn("w:t")))
    normalized = _text(document_text).casefold()
    issues: list[ValidationIssue] = []
    remaining = sorted(set(_TOKEN_PATTERN.findall(document_text)))
    for token in remaining:
        issues.append(
            ValidationIssue(
                "integrity.anchor_remaining",
                "Rendered document still contains an unresolved token anchor.",
                path="document.tokens",
                actual=token,
            )
        )
    for slot in slots:
        semantic = slot["semantic"]
        for marker in _coverage_markers(slot, sources.get(slot["source"])):
            if marker.casefold() not in normalized:
                issues.append(
                    ValidationIssue(
                        "integrity.value_missing",
                        "Fixture value was not found in the rendered document.",
                        semantic,
                        f"semantic.{semantic}",
                        marker,
                        "missing",
                    )
                )
    if manifest.get("legacyRendererUsed") is not False:
        issues.append(
            ValidationIssue(
                "integrity.legacy_renderer",
                "Validation artifact must be produced without the Legacy Renderer.",
                path="manifest.legacyRendererUsed",
                expected=False,
                actual=manifest.get("legacyRendererUsed"),
            )
        )
    return issues


def _coverage_markers(slot: dict[str, Any], value: Any) -> list[str]:
    renderer = slot["renderer"]
    markers: list[str] = []
    if renderer in _TABLE_RENDERERS and isinstance(value, list):
        fields = [item["source"] for item in slot.get("fields", [])]
        if renderer == "finding_table" and not fields:
            fields = ["hostname", "finding", "evidence"]
        for row in value:
            if isinstance(row, dict):
                markers.extend(_text(row.get(field)) for field in fields if _text(row.get(field)))
    elif renderer in {"summary_table", "field_group"} and isinstance(value, dict):
        markers.extend(_text(item) for item in value.values() if _text(item))
    elif renderer == "finding_sections" and isinstance(value, list):
        for row in value:
            if isinstance(row, dict):
                for field_name in ("hostname", "finding", "details", "evidence"):
                    marker = _text(row.get(field_name))
                    if marker:
                        markers.append(marker)
    elif renderer == "response_sections" and isinstance(value, dict):
        for items in value.values():
            if isinstance(items, list):
                markers.extend(_text(item) for item in items if _text(item))
    elif isinstance(value, str):
        markers.extend(line.strip() for line in value.splitlines() if line.strip())
    elif isinstance(value, list):
        markers.extend(_text(item) for item in value if _text(item))
    return list(dict.fromkeys(markers))


def _source_row_count(renderer: str, value: Any) -> int:
    if renderer == "text":
        return 1
    if renderer in _TABLE_RENDERERS:
        return len(value) if isinstance(value, (list, dict)) else 0
    if renderer == "finding_sections":
        return len(value) if isinstance(value, list) else 0
    if renderer == "response_sections" and isinstance(value, dict):
        return sum(len(item) if isinstance(item, list) else 1 for item in value.values())
    if isinstance(value, list):
        return len(value)
    if isinstance(value, str):
        return len([line for line in value.splitlines() if line.strip()]) or 1
    if isinstance(value, dict):
        return len(value)
    return 0


def _relationships(archive: zipfile.ZipFile) -> list[dict[str, str]]:
    name = "word/_rels/document.xml.rels"
    if name not in archive.namelist():
        return []
    root = ElementTree.fromstring(archive.read(name))
    return sorted(
        [
            {
                "type": item.attrib.get("Type", "").rsplit("/", 1)[-1],
                "target": item.attrib.get("Target", ""),
                "mode": item.attrib.get("TargetMode", ""),
            }
            for item in root
        ],
        key=lambda item: (item["type"], item["target"], item["mode"]),
    )


def _progress_adapter(callback: Callable[[int, str], Any] | None) -> Callable[[int, str], None]:
    last = -1

    def emit(percent: int, stage: str) -> None:
        nonlocal last
        bounded = min(100, max(last, int(percent)))
        last = bounded
        if callback is not None:
            callback(bounded, stage)

    return emit


def _check_cancelled(callback: Callable[[], Any] | None) -> None:
    if callback is None:
        return
    try:
        cancelled = callback()
    except ProfileRenderCancelled:
        raise
    except Exception as exc:
        raise ProfileRenderCancelled("Template Pack validation was cancelled.") from exc
    if cancelled is True:
        raise ProfileRenderCancelled("Template Pack validation was cancelled.")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(_text(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_text(item)}" for key, item in value.items())
    return unicodedata.normalize("NFC", str(value)).strip()


def _canonical_sha256(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return _sha256(payload.encode("utf-8"))


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _timestamp() -> str:
    return datetime.now(timezone.utc).isoformat()
