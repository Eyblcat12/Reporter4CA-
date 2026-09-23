"""Loopback-only real API test server; all mutable data lives in a temporary directory."""

from __future__ import annotations

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps" / "backend"))
import uvicorn  # noqa: E402
from core.template_profiles import REPORT_REQUIREMENTS  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.responses import Response  # noqa: E402

from tests.test_api_integration import ApiIntegrationTests  # noqa: E402

if __name__ == "__main__":
    case = ApiIntegrationTests()
    case.setUp()

    @case.client.app.get("/__test/unprepared.docx")
    def unprepared_fixture():
        document = Document()
        document.add_heading("Synthetic customer layout — static heading", 1)
        for semantic in REPORT_REQUIREMENTS["full"]:
            document.add_paragraph("Replace: " + semantic)
        output = io.BytesIO()
        document.save(output)
        return Response(
            output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    try:
        uvicorn.run(case.client.app, host="127.0.0.1", port=8011)
    finally:
        case.tearDown()
