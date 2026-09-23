from __future__ import annotations

import io
import sys
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from tests import test_api_integration
from tests.test_profile_renderer import _published_pack, _template_for

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps" / "backend"))
from core.profile_renderer import render_profile_document  # noqa: E402
from core.template_pack_catalog import build_template_pack  # noqa: E402
from core.template_profiles import COMMON_REQUIREMENTS, REPORT_REQUIREMENTS  # noqa: E402

from tests.test_profile_renderer import _evidence, _prepared  # noqa: E402


class TemplatePackRuntimeTests(unittest.TestCase):
    setUp = test_api_integration.ApiIntegrationTests.setUp
    tearDown = test_api_integration.ApiIntegrationTests.tearDown

    def test_unprepared_docx_normalization_is_copy_only_and_revision_guarded(self):
        document = Document()
        document.add_heading("Customer static heading", 1)
        placements = []
        for semantic in REPORT_REQUIREMENTS["full"]:
            document.add_paragraph(f"Customer placeholder: {semantic}")
            placements.append(
                {
                    "semantic": semantic,
                    "blockIndex": len(document.element.body) - 2,
                    "mode": "replace",
                }
            )
        source = io.BytesIO()
        document.save(source)
        original = self.template_studio.create(
            source.getvalue(),
            report_type="full",
            profile_id="raw-customer",
            display_name="Raw Customer",
        )
        endpoint = f"/api/template-packs/workspaces/{original['workspaceId']}"
        structure = self.client.get(endpoint + "/structure")
        self.assertEqual(structure.status_code, 200)
        request = {"expectedRevision": original["revision"], "placements": placements}
        stale = self.client.post(endpoint + "/normalize", json={**request, "expectedRevision": 99})
        self.assertEqual(stale.status_code, 409)
        incomplete = self.client.post(
            endpoint + "/normalize", json={**request, "placements": placements[:-1]}
        )
        self.assertEqual(incomplete.status_code, 400)
        normalized = self.client.post(endpoint + "/normalize", json=request)
        self.assertEqual(normalized.status_code, 201, normalized.text)
        workspace = normalized.json()
        self.assertNotEqual(workspace["workspaceId"], original["workspaceId"])
        self.assertEqual(
            self.template_studio.source_bytes(original["workspaceId"]), source.getvalue()
        )
        self.assertEqual(workspace["coveragePercent"], 0)
        for semantic in REPORT_REQUIREMENTS["full"]:
            workspace = self.template_studio.approve(
                workspace["workspaceId"],
                semantic=semantic,
                anchor={
                    "kind": "content_control",
                    "value": "REPORTER_" + semantic.upper().replace(".", "_"),
                },
                fields=[
                    {"source": field, "target": f"column:{index}"}
                    for index, field in enumerate(COMMON_REQUIREMENTS[semantic].required_fields, 1)
                ],
                expected_revision=workspace["revision"],
            )
        self.assertEqual(workspace["coveragePercent"], 100)
        normalized_source = self.template_studio.source_bytes(workspace["workspaceId"])
        pack = build_template_pack(workspace, normalized_source, _evidence())
        rendered = render_profile_document(
            pack,
            _prepared("full", normalized_source, {"servers": [], "clients": [], "metadata": {}}),
        )
        xml_text = " ".join(n.text or "" for n in rendered.document.element.body.iter(qn("w:t")))
        self.assertIn("Customer static heading", xml_text)
        self.assertNotIn("Customer placeholder:", xml_text)

    def test_normalized_table_preserves_customer_header_and_cell_format(self):
        from core.template_normalization import normalize_template_source
        from docx.shared import Inches, Pt

        document = Document()
        placements = []
        for semantic in REPORT_REQUIREMENTS["full"]:
            fields = COMMON_REQUIREMENTS[semantic].required_fields
            if semantic == "inventory.server":
                table = document.add_table(rows=2, cols=len(fields))
                table.style = "Table Grid"
                for index, field in enumerate(fields):
                    table.columns[index].width = Inches(1)
                    table.cell(0, index).text = "Customer " + field
                    run = table.cell(1, index).paragraphs[0].add_run("old")
                    run.font.size = Pt(9)
                block = len(document.element.body) - 2
            else:
                document.add_paragraph(semantic)
                block = len(document.element.body) - 2
            placements.append({"semantic": semantic, "blockIndex": block, "mode": "replace"})
        # The exact semantic is read from the canonical catalog, not guessed.
        output = io.BytesIO()
        document.save(output)
        source = normalize_template_source(output.getvalue(), "full", placements)
        anchors = {
            semantic: {
                "kind": "content_control",
                "value": "REPORTER_" + semantic.upper().replace(".", "_"),
            }
            for semantic in REPORT_REQUIREMENTS["full"]
        }
        pack = _published_pack(source, "full", anchors)
        rendered = render_profile_document(
            pack,
            _prepared(
                "full",
                source,
                {
                    "servers": [
                        {
                            "hostname": "STYLE-SRV",
                            "ip": "10.0.0.1",
                            "os": "Linux",
                            "result": "Clean",
                        }
                    ],
                    "clients": [],
                    "metadata": {},
                },
            ),
        )
        text = " ".join(n.text or "" for n in rendered.document.element.body.iter(qn("w:t")))
        self.assertIn("STYLE-SRV", text)
        self.assertIn("Customer hostname", text)
        styled_table = next(
            node
            for node in rendered.document.element.body.iter(qn("w:tbl"))
            if "Customer hostname" in " ".join(t.text or "" for t in node.iter(qn("w:t")))
        )
        self.assertTrue(
            any(node.get(qn("w:val")) == "18" for node in styled_table.iter(qn("w:sz")))
        )

    def wait_job(self, url, *, preview=False):
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, response.text)
            state = response.json() if preview else response.json()["job"]
            if state["status"] in {"ready", "completed", "failed", "cancelled"}:
                self.assertIn(state["status"], {"ready", "completed"}, state)
                return state
            time.sleep(0.02)
        self.fail("Report job timed out")

    def test_three_published_pack_types_preview_generate_and_scope(self):
        for report_type in ("full", "server_only", "client_only"):
            with self.subTest(report_type=report_type):
                template, anchors = _template_for(report_type)
                pack = _published_pack(template, report_type, anchors)
                catalog = self.template_pack_catalog.snapshot()
                installed = self.template_pack_catalog.install(
                    pack, expected_revision=catalog["revision"]
                )
                pack_id = f"renderer-{report_type.replace('_', '-')}"
                version = next(iter(installed["packs"][pack_id]["versions"]))
                selector = f"rptpack:{pack_id}:{version}"
                request = {
                    "rows": [
                        {
                            "type": "server",
                            "hostname": "PACK-SRV-01",
                            "ip": "10.1.1.1",
                            "os": "Linux",
                            "result": "Malware detected",
                        },
                        {
                            "type": "client",
                            "hostname": "PACK-PC-01",
                            "ip": "10.1.1.2",
                            "os": "Windows 11",
                            "result": "Clean",
                        },
                    ],
                    "reportType": report_type,
                    "templatePath": selector,
                    "title": "Pack runtime integration",
                    "disablePlugins": True,
                }
                with patch(
                    "api.routes.generate_report",
                    side_effect=AssertionError("Legacy renderer must not be used"),
                ):
                    listed = self.client.get("/api/templates").json()["templates"]
                    self.assertTrue(
                        any(t["path"] == selector and not t["isDefault"] for t in listed)
                    )
                    preview = self.client.post("/api/preview-jobs", json=request)
                    self.assertEqual(preview.status_code, 202, preview.text)
                    preview_id = preview.json()["previewId"]
                    self.wait_job(f"/api/preview-jobs/{preview_id}", preview=True)
                    content = self.client.get(f"/api/preview-jobs/{preview_id}/content")
                    self.assertEqual(content.status_code, 200)
                    text = " ".join(
                        n.text or ""
                        for n in Document(io.BytesIO(content.content)).element.body.iter(qn("w:t"))
                    )
                    self.assertIn("Pack runtime integration", text)
                    self.assertNotIn("{{", text)
                    for hostname, included in (
                        ("PACK-SRV-01", report_type != "client_only"),
                        ("PACK-PC-01", report_type != "server_only"),
                    ):
                        self.assertEqual(hostname in text, included, text)
                    generated = self.client.post(
                        "/api/report-jobs", json={**request, "previewId": preview_id}
                    )
                    self.assertEqual(generated.status_code, 202, generated.text)
                    job_id = generated.json()["job"]["id"]
                    self.wait_job(f"/api/report-jobs/{job_id}")
                    download = self.client.get(f"/api/report-jobs/{job_id}/download")
                    self.assertEqual(download.status_code, 200)
                    self.assertEqual(download.content, content.content)
                    # Cold generation uses the same profile dispatch, without the cache.
                    cold = self.client.post(
                        "/api/report-jobs", json={**request, "title": "Cold profile generation"}
                    )
                    self.assertEqual(cold.status_code, 202, cold.text)
                    self.wait_job(f"/api/report-jobs/{cold.json()['job']['id']}")
                with patch("api.routes.template_packs_enabled", return_value=False):
                    blocked = self.client.post("/api/report-jobs", json=request)
                    self.assertEqual(blocked.status_code, 409, blocked.text)
                wrong_type = self.client.post(
                    "/api/report-jobs", json={**request, "reportType": "technical"}
                )
                self.assertEqual(wrong_type.status_code, 422, wrong_type.text)


if __name__ == "__main__":
    unittest.main()
