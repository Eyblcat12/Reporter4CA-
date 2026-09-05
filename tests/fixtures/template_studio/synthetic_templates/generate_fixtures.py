"""Generate structurally distinct DOCX fixtures for Template Studio pilots.

These documents contain no customer data. They deliberately exercise token,
bookmark, and content-control anchors while keeping the visible layout suitable
for human visual review.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
FIXTURES = {
    "full": {
        "filename": "cross_platform_assessment_full.docx",
        "title": "CROSS-PLATFORM SECURITY ASSESSMENT",
        "accent": "1F4E78",
        "anchors": {
            "report.title": {"kind": "bookmark", "value": "report_title"},
            "overview": {"kind": "content_control", "value": "overview"},
            "inventory.server": {"kind": "token", "value": "{{SERVER_INVENTORY}}"},
            "inventory.client": {"kind": "token", "value": "{{CLIENT_INVENTORY}}"},
            "results.server": {"kind": "content_control", "value": "server_results"},
            "results.client": {"kind": "content_control", "value": "client_results"},
            "investigation": {"kind": "bookmark", "value": "investigation"},
            "remediation": {"kind": "token", "value": "{{REMEDIATION_REGISTER}}"},
            "ioc": {"kind": "content_control", "value": "ioc_register"},
            "recommendations": {"kind": "bookmark", "value": "recommendations"},
        },
    },
    "server_only": {
        "filename": "infrastructure_review_server.docx",
        "title": "INFRASTRUCTURE SECURITY REVIEW",
        "accent": "245B47",
        "anchors": {
            "report.title": {"kind": "content_control", "value": "report_title"},
            "overview": {"kind": "bookmark", "value": "overview"},
            "inventory.server": {"kind": "content_control", "value": "server_inventory"},
            "results.server": {"kind": "token", "value": "{{SERVER_RESULTS}}"},
            "investigation": {"kind": "content_control", "value": "investigation"},
            "remediation": {"kind": "bookmark", "value": "remediation"},
            "ioc": {"kind": "token", "value": "{{IOC_REGISTER}}"},
            "recommendations": {"kind": "content_control", "value": "recommendations"},
        },
    },
    "client_only": {
        "filename": "endpoint_assurance_client.docx",
        "title": "ENDPOINT ASSURANCE REPORT",
        "accent": "6B3FA0",
        "anchors": {
            "report.title": {"kind": "token", "value": "{{REPORT_TITLE}}"},
            "overview": {"kind": "content_control", "value": "overview"},
            "inventory.client": {"kind": "bookmark", "value": "client_inventory"},
            "results.client": {"kind": "content_control", "value": "client_results"},
            "investigation": {"kind": "token", "value": "{{INVESTIGATION}}"},
            "remediation": {"kind": "content_control", "value": "remediation"},
            "ioc": {"kind": "bookmark", "value": "ioc_register"},
            "recommendations": {"kind": "token", "value": "{{RECOMMENDATIONS}}"},
        },
    },
}


def _setup(document: Document, accent: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("24313D")
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        borders = style._element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            style._element.get_or_add_pPr().remove(borders)
    header = section.header.paragraphs[0]
    header.text = "REPORTER PRO  /  CONTROLLED ASSESSMENT"
    header.style = document.styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(accent)
    footer = section.footer.paragraphs[0]
    footer.text = "Confidential — generated from a validated Template Pack"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), fill)


def _summary_band(document: Document, accent: str, scope: str) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, heading, value in zip(
        table.rows[0].cells,
        ("SCOPE", "SOURCE", "OUTPUT"),
        (scope, "Tracking dataset", "Validated DOCX"),
        strict=True,
    ):
        _shade(cell, "EAF0F5")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{heading}\n{value}")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(accent)
    document.add_paragraph()


def _bookmark(document: Document, name: str, text: str, bookmark_id: int) -> None:
    paragraph = document.add_paragraph()
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(start)
    paragraph.add_run(text)
    paragraph._p.append(end)


def _content_control(document: Document, tag_value: str, text: str) -> None:
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), tag_value)
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), tag_value.replace("_", " ").title())
    properties.extend((alias, tag))
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph.append(run)
    content.append(paragraph)
    control.extend((properties, content))
    document.element.body.insert(len(document.element.body) - 1, control)


def _anchor(document: Document, anchor: dict[str, str], placeholder: str, bookmark_id: int) -> int:
    if anchor["kind"] == "token":
        paragraph = document.add_paragraph(anchor["value"])
        paragraph.paragraph_format.keep_together = True
    elif anchor["kind"] == "bookmark":
        _bookmark(document, anchor["value"], placeholder, bookmark_id)
        bookmark_id += 1
    else:
        _content_control(document, anchor["value"], placeholder)
    return bookmark_id


def _cover(document: Document, title: str, scope: str) -> None:
    document.add_paragraph("SECURITY OPERATIONS", style="Subtitle")
    paragraph = document.add_paragraph(title, style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.add_paragraph(
        "Evidence-led assessment of assets, findings, indicators and remediation actions."
    )
    _summary_band(document, document.styles["Title"].font.color.rgb.__str__(), scope)


def _build(report_type: str, config: dict) -> Document:
    document = Document()
    _setup(document, config["accent"])
    _cover(document, config["title"], report_type.replace("_", " ").title())
    bookmark_id = 1
    anchors = config["anchors"]
    bookmark_id = _anchor(document, anchors["report.title"], "Report title", bookmark_id)

    sections: list[tuple[str, str]] = [("1. Executive context", "overview")]
    if report_type in {"full", "server_only"}:
        sections.extend(
            [
                ("2. Server estate", "inventory.server"),
                ("3. Server assessment", "results.server"),
            ]
        )
    if report_type in {"full", "client_only"}:
        sections.extend(
            [
                (
                    "4. Endpoint estate" if report_type == "full" else "2. Endpoint estate",
                    "inventory.client",
                ),
                (
                    "5. Endpoint assessment" if report_type == "full" else "3. Endpoint assessment",
                    "results.client",
                ),
            ]
        )
    tail_start = 6 if report_type == "full" else 4
    sections.extend(
        [
            (f"{tail_start}. Investigation workspace", "investigation"),
            (f"{tail_start + 1}. Remediation register", "remediation"),
            (f"{tail_start + 2}. Indicator register", "ioc"),
            (f"{tail_start + 3}. Recommendations", "recommendations"),
        ]
    )
    for heading, semantic in sections:
        if report_type == "full" and semantic == "ioc":
            document.add_page_break()
        document.add_heading(heading, level=1)
        document.add_paragraph(
            "This section is populated from the normalized, evidence-traceable report snapshot.",
            style="Caption",
        )
        bookmark_id = _anchor(
            document,
            anchors[semantic],
            f"Mapped content: {semantic}",
            bookmark_id,
        )

    if report_type == "client_only":
        section = document.add_section()
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        document.add_heading("Appendix — Review notes", level=1)
        table = document.add_table(rows=3, cols=4)
        table.style = "Table Grid"
        for index, text in enumerate(("Owner", "Status", "Evidence", "Decision")):
            table.cell(0, index).text = text
            _shade(table.cell(0, index), "EDE5F5")
        for row in table.rows[1:]:
            for cell in row.cells:
                cell.text = "Reserved for analyst review"
    return document


def main() -> None:
    manifest = {"schemaVersion": 1, "fixtures": {}}
    for report_type, config in FIXTURES.items():
        path = ROOT / config["filename"]
        document = _build(report_type, config)
        document.core_properties.title = config["title"]
        document.core_properties.subject = "Reporter Pro Template Studio synthetic pilot fixture"
        document.core_properties.author = "Reporter Pro"
        document.save(path)
        payload = path.read_bytes()
        manifest["fixtures"][report_type] = {
            "file": path.name,
            "sha256": hashlib.sha256(payload).hexdigest(),
            "sizeBytes": len(payload),
            "anchors": config["anchors"],
        }
    (ROOT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
