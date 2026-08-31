from __future__ import annotations

import ast
import hashlib
import io
import json
import os
import random
import stat
import unittest
import warnings
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apps.backend.core.config import template_packs_enabled
from apps.backend.core.template_pack import TemplatePackError, inspect_template_pack
from apps.backend.core.template_profile_analyzer import analyze_profile_template
from apps.backend.core.template_profiles import (
    COMMON_REQUIREMENTS,
    PROFILE_SCHEMA_VERSION,
    REPORT_REQUIREMENTS,
    validate_template_profile,
)

ROOT = Path(__file__).resolve().parents[1]


def _complete_profile(report_type: str = "full") -> dict:
    slots = []
    for index, semantic in enumerate(REPORT_REQUIREMENTS[report_type]):
        requirement = COMMON_REQUIREMENTS[semantic]
        slots.append(
            {
                "id": semantic.replace(".", "_"),
                "semantic": semantic,
                "renderer": requirement.renderer,
                "source": requirement.source,
                "anchor": {
                    "kind": "content_control",
                    "value": f"REPORTER_SLOT_{index}",
                },
                "repeat": "once",
                "fields": [
                    {"source": field, "target": f"column:{field}"}
                    for field in requirement.required_fields
                ],
            }
        )
    return {
        "schemaVersion": PROFILE_SCHEMA_VERSION,
        "profileId": "customer-security-report",
        "displayName": "Customer Security Report",
        "version": "1.0.0",
        "reportType": report_type,
        "status": "published",
        "templateSha256": "0" * 64,
        "slots": slots,
        "validation": {
            "fixturePassed": True,
            "integrityPassed": True,
            "visualApproved": True,
        },
    }


def _minimal_docx() -> bytes:
    document = Document()
    document.add_heading("Customer report", level=1)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _analyzable_docx() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Customer header"
    document.sections[0].footer.paragraphs[0].text = "Customer footer"
    document.add_heading("Overview", level=1)
    token_paragraph = document.add_paragraph()
    token_paragraph.add_run("{{REPORT_")
    token_paragraph.add_run("TITLE}}")
    document.add_paragraph("{{REPORT_TITLE}}")

    bookmark_paragraph = document.add_paragraph("Indicators")
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "42")
    bookmark_start.set(qn("w:name"), "REPORTER_IOC")
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "42")
    bookmark_paragraph._p.insert(0, bookmark_start)
    bookmark_paragraph._p.append(bookmark_end)

    content_control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "REPORTER_OVERVIEW")
    properties.append(tag)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    content.append(paragraph)
    content_control.append(properties)
    content_control.append(content)
    document._element.body.insert(0, content_control)

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Hostname"
    table.rows[0].cells[1].text = "Result"
    document.add_page_break()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _canonical(value: dict) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def _pack_bytes(*, omit_last_slot: bool = False, checksum_mismatch: bool = False) -> bytes:
    profile = _complete_profile()
    template = _minimal_docx()
    layout_slots = []
    mapping_slots = []
    source_slots = profile["slots"][:-1] if omit_last_slot else profile["slots"]
    for slot in source_slots:
        layout_slots.append(
            {key: slot[key] for key in ("id", "semantic", "renderer", "anchor", "repeat")}
        )
        mapping_slots.append({"id": slot["id"], "source": slot["source"], "fields": slot["fields"]})
    layout = {
        key: profile[key]
        for key in (
            "schemaVersion",
            "profileId",
            "displayName",
            "version",
            "reportType",
            "status",
        )
    }
    layout["slots"] = layout_slots
    mapping = {"profileId": profile["profileId"], "slots": mapping_slots}
    validation = profile["validation"]
    payloads = {
        "template.docx": template,
        "layout.json": _canonical(layout),
        "mapping.json": _canonical(mapping),
        "validation.json": _canonical(validation),
    }
    manifest = {
        "formatVersion": "1.0",
        "packId": profile["profileId"],
        "version": profile["version"],
        "checksums": {
            name: hashlib.sha256(payload).hexdigest() for name, payload in payloads.items()
        },
    }
    if checksum_mismatch:
        manifest["checksums"]["layout.json"] = "f" * 64

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", _canonical(manifest))
        for name, payload in payloads.items():
            archive.writestr(name, payload)
    return output.getvalue()


def _replace_pack_template(source: bytes, template: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(source)) as existing:
        values = {info.filename: existing.read(info.filename) for info in existing.infolist()}
    manifest = json.loads(values["manifest.json"])
    manifest["checksums"]["template.docx"] = hashlib.sha256(template).hexdigest()
    values["manifest.json"] = _canonical(manifest)
    values["template.docx"] = template
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in values.items():
            archive.writestr(name, payload)
    return output.getvalue()


class TemplateProfileTests(unittest.TestCase):
    def test_complete_profile_is_publishable(self) -> None:
        result = validate_template_profile(_complete_profile())

        self.assertTrue(result.structural_valid)
        self.assertTrue(result.mapping_complete)
        self.assertEqual(result.coverage_percent, 100.0)
        self.assertTrue(result.publishable)

    def test_missing_semantic_cannot_reach_complete_mapping(self) -> None:
        profile = _complete_profile()
        profile["slots"].pop()

        result = validate_template_profile(profile)

        self.assertTrue(result.structural_valid)
        self.assertFalse(result.mapping_complete)
        self.assertLess(result.coverage_percent, 100.0)
        self.assertFalse(result.publishable)
        self.assertEqual(result.missing_semantics, ("recommendations",))

    def test_duplicate_anchor_is_rejected(self) -> None:
        profile = _complete_profile()
        profile["slots"][1]["anchor"] = dict(profile["slots"][0]["anchor"])

        result = validate_template_profile(profile)

        self.assertFalse(result.structural_valid)
        self.assertFalse(result.publishable)
        self.assertIn("anchor.duplicate", {issue.code for issue in result.errors})

    def test_required_table_field_mapping_is_enforced(self) -> None:
        profile = _complete_profile()
        inventory = next(
            slot for slot in profile["slots"] if slot["semantic"] == "inventory.server"
        )
        inventory["fields"] = [
            field for field in inventory["fields"] if field["source"] != "hostname"
        ]

        result = validate_template_profile(profile)

        self.assertFalse(result.structural_valid)
        self.assertFalse(result.mapping_complete)
        self.assertIn("fields.required_missing", {issue.code for issue in result.errors})

    def test_validation_evidence_is_required_for_publication(self) -> None:
        profile = _complete_profile()
        profile["validation"]["visualApproved"] = False

        result = validate_template_profile(profile)

        self.assertTrue(result.structural_valid)
        self.assertTrue(result.mapping_complete)
        self.assertFalse(result.publishable)
        self.assertIn("validation.incomplete", {issue.code for issue in result.warnings})

    def test_non_published_profile_remains_non_activatable(self) -> None:
        profile = _complete_profile()
        profile["status"] = "test_passed"

        result = validate_template_profile(profile)

        self.assertTrue(result.mapping_complete)
        self.assertFalse(result.publishable)

    def test_template_pack_feature_is_disabled_by_default(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            self.assertFalse(template_packs_enabled())
        with patch.dict(os.environ, {"AUTO_REPORT_TEMPLATE_PACKS": "1"}, clear=True):
            self.assertTrue(template_packs_enabled())


class TemplatePackInspectionTests(unittest.TestCase):
    def test_valid_published_pack_is_accepted_without_installing(self) -> None:
        result = inspect_template_pack(_pack_bytes())

        self.assertEqual(result.pack_id, "customer-security-report")
        self.assertTrue(result.publishable)
        self.assertEqual(result.validation.coverage_percent, 100.0)

    def test_checksum_mismatch_is_rejected(self) -> None:
        with self.assertRaisesRegex(TemplatePackError, "Checksum mismatch"):
            inspect_template_pack(_pack_bytes(checksum_mismatch=True))

    def test_manifest_unknown_fields_are_rejected(self) -> None:
        source = _pack_bytes()
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as existing:
            values = {info.filename: existing.read(info.filename) for info in existing.infolist()}
        manifest = json.loads(values["manifest.json"])
        manifest["executable"] = "payload.py"
        values["manifest.json"] = _canonical(manifest)
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as modified:
            for name, payload in values.items():
                modified.writestr(name, payload)

        with self.assertRaisesRegex(TemplatePackError, "unsupported fields"):
            inspect_template_pack(output.getvalue())

    def test_incomplete_mapping_is_visible_but_not_publishable(self) -> None:
        data = _pack_bytes(omit_last_slot=True)

        inspection = inspect_template_pack(data, require_publishable=False)

        self.assertFalse(inspection.publishable)
        self.assertEqual(inspection.validation.missing_semantics, ("recommendations",))
        with self.assertRaisesRegex(TemplatePackError, "not publishable"):
            inspect_template_pack(data)

    def test_path_traversal_is_rejected_before_extraction(self) -> None:
        source = _pack_bytes()
        output = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(source)) as existing,
            zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as modified,
        ):
            for info in existing.infolist():
                modified.writestr(info.filename, existing.read(info.filename))
            modified.writestr("../payload.py", b"print('unsafe')")

        with self.assertRaisesRegex(TemplatePackError, "Unsafe archive member"):
            inspect_template_pack(output.getvalue())

    def test_invalid_docx_package_is_rejected(self) -> None:
        source = _pack_bytes()
        output = io.BytesIO()
        fake_docx = b"PK\x03\x04not-a-real-office-package"
        with zipfile.ZipFile(io.BytesIO(source)) as existing:
            values = {info.filename: existing.read(info.filename) for info in existing.infolist()}
        manifest = json.loads(values["manifest.json"])
        manifest["checksums"]["template.docx"] = hashlib.sha256(fake_docx).hexdigest()
        values["manifest.json"] = _canonical(manifest)
        values["template.docx"] = fake_docx
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as modified:
            for name, payload in values.items():
                modified.writestr(name, payload)

        with self.assertRaises(TemplatePackError):
            inspect_template_pack(output.getvalue())

    def test_duplicate_member_and_symbolic_link_are_rejected(self) -> None:
        source = _pack_bytes()
        with zipfile.ZipFile(io.BytesIO(source)) as existing:
            values = [(info.filename, existing.read(info.filename)) for info in existing.infolist()]

        duplicate = io.BytesIO()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(duplicate, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, payload in values:
                    archive.writestr(name, payload)
                archive.writestr("manifest.json", values[0][1])
        with self.assertRaisesRegex(TemplatePackError, "Duplicate archive member"):
            inspect_template_pack(duplicate.getvalue())

        linked = io.BytesIO()
        with zipfile.ZipFile(linked, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, payload in values:
                if name == "mapping.json":
                    info = zipfile.ZipInfo(name)
                    info.create_system = 3
                    info.external_attr = (stat.S_IFLNK | 0o777) << 16
                    archive.writestr(info, payload)
                else:
                    archive.writestr(name, payload)
        with self.assertRaisesRegex(TemplatePackError, "Symbolic links"):
            inspect_template_pack(linked.getvalue())

    def test_compression_bomb_metadata_is_rejected_before_json_processing(self) -> None:
        source = _pack_bytes()
        with zipfile.ZipFile(io.BytesIO(source)) as existing:
            values = {info.filename: existing.read(info.filename) for info in existing.infolist()}
        values["layout.json"] = b"0" * (1024 * 1024)
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for name, payload in values.items():
                archive.writestr(name, payload)

        with self.assertRaisesRegex(TemplatePackError, "compression ratio"):
            inspect_template_pack(output.getvalue())

    def test_deeply_nested_json_returns_a_bounded_pack_error(self) -> None:
        source = _pack_bytes()
        with zipfile.ZipFile(io.BytesIO(source)) as existing:
            values = {info.filename: existing.read(info.filename) for info in existing.infolist()}
        values["layout.json"] = b'{"displayName":' + b"[" * 2000 + b"0" + b"]" * 2000 + b"}"
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_STORED) as archive:
            for name, payload in values.items():
                archive.writestr(name, payload)

        with self.assertRaisesRegex(TemplatePackError, "nesting exceeds"):
            inspect_template_pack(output.getvalue())

    def test_unsafe_nested_ooxml_and_inner_compression_bomb_are_rejected(self) -> None:
        template = _minimal_docx()
        malicious = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(template)) as existing,
            zipfile.ZipFile(malicious, "w", zipfile.ZIP_DEFLATED) as modified,
        ):
            for info in existing.infolist():
                payload = existing.read(info.filename)
                if info.filename == "word/document.xml":
                    payload = b"<!DOCTYPE w:document [<!ENTITY x 'unsafe'>]>" + payload
                modified.writestr(info.filename, payload)
        with self.assertRaisesRegex(TemplatePackError, "Unsafe XML declaration"):
            inspect_template_pack(_replace_pack_template(_pack_bytes(), malicious.getvalue()))

        compressed = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(template)) as existing,
            zipfile.ZipFile(compressed, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as modified,
        ):
            for info in existing.infolist():
                payload = existing.read(info.filename)
                if info.filename == "word/document.xml":
                    payload = b"0" * (1024 * 1024)
                modified.writestr(info.filename, payload)
        with self.assertRaisesRegex(TemplatePackError, "DOCX compression ratio"):
            inspect_template_pack(_replace_pack_template(_pack_bytes(), compressed.getvalue()))

    def test_seeded_mutation_corpus_never_leaks_parser_exceptions(self) -> None:
        source = _pack_bytes()
        randomizer = random.Random(0xC0DEC0DE)
        for case in range(128):
            mutated = bytearray(source)
            operation = case % 4
            if operation == 0:
                del mutated[randomizer.randrange(2, len(mutated)) :]
            elif operation == 1:
                for _ in range(randomizer.randint(1, 4)):
                    index = randomizer.randrange(len(mutated))
                    mutated[index] ^= randomizer.randrange(1, 256)
            elif operation == 2:
                mutated.extend(randomizer.randbytes(randomizer.randint(1, 32)))
            else:
                start = randomizer.randrange(len(mutated) - 1)
                length = randomizer.randint(1, min(32, len(mutated) - start))
                mutated[start : start + length] = b"\0" * length
            with self.subTest(case=case, operation=operation):
                try:
                    inspect_template_pack(bytes(mutated))
                except TemplatePackError:
                    pass

    def test_deleted_zip_range_never_leaks_negative_seek_value_error(self) -> None:
        source = _pack_bytes()
        mutated = bytearray(source)
        del mutated[11201:11217]

        with self.assertRaisesRegex(TemplatePackError, "ZIP is invalid"):
            inspect_template_pack(bytes(mutated))


class TemplateProfileAnalyzerTests(unittest.TestCase):
    def test_analyzer_reports_facts_but_never_auto_approves_mapping(self) -> None:
        result = analyze_profile_template(_analyzable_docx(), "full")

        self.assertEqual(result["mapping"]["coveragePercent"], 0.0)
        self.assertEqual(result["mapping"]["approvedCount"], 0)
        self.assertEqual(result["profileSeed"]["status"], "analyzed")
        self.assertEqual(result["profileSeed"]["slots"], [])
        self.assertIn("{{REPORT_TITLE}}", result["facts"]["tokens"])
        self.assertEqual(result["facts"]["contentControls"][0]["value"], "REPORTER_OVERVIEW")
        self.assertEqual(result["facts"]["bookmarks"][0]["value"], "REPORTER_IOC")
        self.assertEqual(result["facts"]["headings"][0]["text"], "Overview")
        self.assertEqual(result["facts"]["tables"][0]["headers"], ["Hostname", "Result"])
        self.assertEqual(result["facts"]["tokenOccurrences"]["{{REPORT_TITLE}}"], 2)
        self.assertTrue(result["facts"]["anchorConflicts"])
        self.assertEqual(len(result["facts"]["sections"]), 1)
        self.assertTrue(result["facts"]["headers"][0]["hasText"])
        self.assertTrue(result["facts"]["footers"][0]["hasText"])
        self.assertEqual(result["facts"]["pageBreakCount"], 1)
        self.assertIn("instanceCount", result["facts"]["numbering"])

        checklist = {item["semantic"]: item for item in result["mapping"]["checklist"]}
        self.assertEqual(checklist["overview"]["suggestedAnchors"][0]["approval"], "required")
        self.assertEqual(
            checklist["report.title"]["suggestedAnchors"][0]["confidence"], "exact_name"
        )
        self.assertEqual(checklist["ioc"]["suggestedAnchors"][0]["kind"], "bookmark")

    def test_analyzer_rejects_unknown_report_type(self) -> None:
        with self.assertRaisesRegex(ValueError, "Unsupported report type"):
            analyze_profile_template(_minimal_docx(), "customer_special")

    def test_analyzer_rejects_xml_entity_declarations(self) -> None:
        source = _minimal_docx()
        output = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(source)) as existing,
            zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as modified,
        ):
            for info in existing.infolist():
                payload = existing.read(info.filename)
                if info.filename == "word/document.xml":
                    payload = b"<!DOCTYPE w:document [<!ENTITY x 'unsafe'>]>" + payload
                modified.writestr(info.filename, payload)

        with self.assertRaisesRegex(ValueError, "declarations are not allowed"):
            analyze_profile_template(output.getvalue(), "full")


class TemplatePackArchitectureTests(unittest.TestCase):
    def test_committed_json_schemas_are_valid_and_versioned(self) -> None:
        schema_root = ROOT / "apps" / "backend" / "config" / "template_packs"
        expected = {
            "manifest.schema.json",
            "layout.schema.json",
            "mapping.schema.json",
            "validation.schema.json",
        }

        self.assertEqual({path.name for path in schema_root.glob("*.json")}, expected)
        for name in expected:
            schema = json.loads((schema_root / name).read_text(encoding="utf-8"))
            self.assertEqual(schema["$schema"], "https://json-schema.org/draft/2020-12/schema")
            self.assertIn("1.0", schema["$id"])
            self.assertEqual(schema["type"], "object")

    def test_legacy_generation_modules_do_not_import_template_pack_runtime(self) -> None:
        forbidden = {
            "core.profile_renderer",
            "core.template_pack",
            "core.template_pack_catalog",
            "core.template_pack_preview",
            "core.template_pack_validation",
            "core.template_mapping_workspace",
            "core.template_profile_analyzer",
            "core.template_profiles",
            "core.template_workspace_retention",
            "core.template_workspace_transfer",
        }
        legacy_modules = (
            ROOT / "apps" / "backend" / "core" / "report_generator.py",
            ROOT / "apps" / "backend" / "core" / "report_orchestrator.py",
            ROOT / "apps" / "backend" / "core" / "report_snapshot.py",
        )

        for path in legacy_modules:
            tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
            imports = {
                node.module
                for node in ast.walk(tree)
                if isinstance(node, ast.ImportFrom) and node.module
            }
            imports.update(
                alias.name
                for node in ast.walk(tree)
                if isinstance(node, ast.Import)
                for alias in node.names
            )
            self.assertFalse(forbidden & imports, f"Legacy boundary violated by {path.name}")


if __name__ == "__main__":
    unittest.main()
