from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "apps" / "backend"
sys.path.insert(0, str(BACKEND))

from core.profile_renderer import (  # noqa: E402
    ProfileRenderCancelled,
    ProfileRendererError,
    render_profile_document,
    save_profile_report_atomic,
)
from core.report_snapshot import AcceptedReportSnapshot, PreparedReportSnapshot  # noqa: E402
from core.template_mapping_workspace import TemplateStudioService  # noqa: E402
from core.template_pack_catalog import (  # noqa: E402
    TemplatePackValidationEvidence,
    build_template_pack,
)
from core.template_profiles import COMMON_REQUIREMENTS, REPORT_REQUIREMENTS  # noqa: E402


class ProfileRendererTests(unittest.TestCase):
    def test_all_six_report_types_render_without_legacy_fallback(self) -> None:
        for report_type in REPORT_REQUIREMENTS:
            with self.subTest(report_type=report_type):
                template_bytes, anchors = _template_for(report_type)
                pack = _published_pack(template_bytes, report_type, anchors)
                prepared = _prepared(
                    report_type,
                    template_bytes,
                    {"servers": [], "clients": [], "metadata": {}},
                )

                result = render_profile_document(pack, prepared)

                self.assertEqual(
                    result.manifest["renderedSemantics"],
                    list(REPORT_REQUIREMENTS[report_type]),
                )
                self.assertFalse(result.manifest["legacyRendererUsed"])

    def test_summary_pack_renders_token_bookmark_and_content_control_anchors(self) -> None:
        template_bytes, anchors = _template_for("summary", mixed_anchors=True)
        pack = _published_pack(template_bytes, "summary", anchors)
        prepared = _prepared(
            "summary",
            template_bytes,
            {
                "servers": [],
                "clients": [
                    {
                        "hostname": "PC-001",
                        "ip": "10.0.0.10",
                        "os": "Windows 11",
                        "result": "Detected suspicious activity",
                        "assessment": {
                            "classification": "anomaly",
                            "label": "Ghi nhận dấu hiệu bất thường",
                        },
                        "findings": [
                            {
                                "name": "Proxy tool detected",
                                "severity": "medium",
                                "classification": "anomaly",
                                "evidence": [{"field": "notes", "value": "ngrok.exe"}],
                                "remediation": "Validate and remove unauthorized tooling.",
                            }
                        ],
                    }
                ],
                "metadata": {"recommendations": ["Review the affected endpoint."]},
            },
        )
        progress: list[tuple[int, str]] = []

        result = render_profile_document(
            pack,
            prepared,
            on_progress=lambda percent, semantic: progress.append((percent, semantic)),
        )
        output = io.BytesIO()
        result.document.save(output)
        rendered = Document(io.BytesIO(output.getvalue()))
        xml_text = " ".join(node.text or "" for node in rendered.element.body.iter(qn("w:t")))

        self.assertIn("Profile Renderer Test", xml_text)
        self.assertIn("Proxy tool detected", xml_text)
        self.assertIn("Review the affected endpoint", xml_text)
        self.assertNotIn("{{", xml_text)
        self.assertGreaterEqual(len(rendered.tables), 2)
        self.assertEqual(progress[-1][0], 100)
        self.assertEqual(
            result.manifest["renderedSemantics"],
            list(REPORT_REQUIREMENTS["summary"]),
        )
        self.assertFalse(result.manifest["legacyRendererUsed"])

    def test_server_pack_renders_assets_results_findings_remediation_and_iocs(self) -> None:
        template_bytes, anchors = _template_for("server_only")
        pack = _published_pack(template_bytes, "server_only", anchors)
        prepared = _prepared(
            "server_only",
            template_bytes,
            {
                "servers": [
                    {
                        "hostname": "SRV-001",
                        "ip": "10.0.0.1",
                        "os": "Windows Server 2022",
                        "result": "Malware detected",
                        "assessment": {
                            "classification": "anomaly",
                            "label": "Ghi nhận dấu hiệu bất thường",
                        },
                        "findings": [
                            {
                                "name": "Malware evidence",
                                "severity": "high",
                                "classification": "anomaly",
                                "evidence": [{"field": "result", "value": "trojan"}],
                                "remediation": "Isolate and investigate the server.",
                            }
                        ],
                        "iocs": [{"type": "hash", "value": "a" * 64}],
                    }
                ],
                "clients": [],
                "metadata": {},
            },
        )

        result = render_profile_document(pack, prepared)
        output = io.BytesIO()
        result.document.save(output)
        rendered = Document(io.BytesIO(output.getvalue()))
        text = " ".join(
            cell.text for table in rendered.tables for row in table.rows for cell in row.cells
        )

        self.assertIn("SRV-001", text)
        self.assertIn("Malware detected", text)
        self.assertIn("Ghi nhận dấu hiệu bất thường", text)
        self.assertIn("a" * 64, text)
        self.assertGreaterEqual(len(rendered.tables), 4)

    def test_missing_anchor_fails_with_semantic_context_and_never_falls_back(self) -> None:
        template_bytes, anchors = _template_for("summary")
        with tempfile.TemporaryDirectory() as temporary:
            service = TemplateStudioService(Path(temporary) / "studio")
            workspace = _complete_workspace(service, template_bytes, "summary", anchors)
            workspace["slots"][0]["anchor"]["value"] = "{{MISSING_ANCHOR}}"
            pack = build_template_pack(workspace, template_bytes, _evidence())
        prepared = _prepared("summary", template_bytes, {"servers": [], "clients": []})

        with self.assertRaises(ProfileRendererError) as captured:
            render_profile_document(pack, prepared)

        self.assertEqual(captured.exception.semantic, "report.title")
        self.assertEqual(captured.exception.anchor_value, "{{MISSING_ANCHOR}}")
        self.assertIn("exactly once", str(captured.exception))

    def test_incident_pack_renders_traceability_and_response_blocks(self) -> None:
        template_bytes, anchors = _template_for("incident_response")
        pack = _published_pack(template_bytes, "incident_response", anchors)
        prepared = _prepared(
            "incident_response",
            template_bytes,
            {
                "servers": [],
                "clients": [],
                "metadata": {
                    "incident": {
                        "incident_id": "IR-2026-001",
                        "severity": "high",
                        "status": "contained",
                        "detected_at": "2026-08-28T08:00:00Z",
                        "executive_summary": "Evidence-backed incident summary.",
                        "affected_assets": [{"hostname": "SRV-IR", "ip": "10.1.1.1"}],
                        "timeline": [
                            {
                                "timestamp": "2026-08-28T08:00:00Z",
                                "event": "Alert received",
                                "evidence": "EDR-42",
                            }
                        ],
                        "iocs": [{"type": "domain", "value": "example.invalid"}],
                        "mitre": [
                            {
                                "technique_id": "T1059",
                                "technique": "Command and Scripting Interpreter",
                                "tactic": "Execution",
                            }
                        ],
                        "containment": ["Isolated affected host"],
                        "eradication": ["Removed malicious persistence"],
                        "recovery": ["Restored service and monitored"],
                        "lessons_learned": "Improve alert triage coverage.",
                    }
                },
            },
        )

        result = render_profile_document(pack, prepared)
        output = io.BytesIO()
        result.document.save(output)
        rendered = Document(io.BytesIO(output.getvalue()))
        xml_text = " ".join(node.text or "" for node in rendered.element.body.iter(qn("w:t")))

        for expected in (
            "IR-2026-001",
            "Alert received",
            "T1059",
            "example.invalid",
            "Isolated affected host",
            "Improve alert triage coverage",
        ):
            self.assertIn(expected, xml_text)
        self.assertEqual(
            result.manifest["renderedSemantics"],
            list(REPORT_REQUIREMENTS["incident_response"]),
        )

    def test_report_type_and_template_snapshot_mismatch_are_rejected(self) -> None:
        template_bytes, anchors = _template_for("summary")
        pack = _published_pack(template_bytes, "summary", anchors)
        wrong_type = _prepared("technical", template_bytes, {"servers": [], "clients": []})
        other_template, _ = _template_for("summary")
        mismatched_template = _prepared("summary", other_template + b"changed", {})

        with self.assertRaisesRegex(ProfileRendererError, "report type"):
            render_profile_document(pack, wrong_type)
        with self.assertRaisesRegex(ProfileRendererError, "pinned"):
            render_profile_document(pack, mismatched_template)

    def test_render_honors_cooperative_cancellation(self) -> None:
        template_bytes, anchors = _template_for("summary")
        pack = _published_pack(template_bytes, "summary", anchors)
        prepared = _prepared("summary", template_bytes, {"servers": [], "clients": []})

        with self.assertRaises(ProfileRenderCancelled):
            render_profile_document(pack, prepared, check_cancelled=lambda: True)

    def test_cancellation_removes_temporary_output(self) -> None:
        template_bytes, anchors = _template_for("summary")
        pack = _published_pack(template_bytes, "summary", anchors)
        prepared = _prepared("summary", template_bytes, {"servers": [], "clients": []})
        result = render_profile_document(pack, prepared)
        calls = 0

        def cancel_after_save() -> bool:
            nonlocal calls
            calls += 1
            return calls >= 2

        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "profile-report.docx"
            with self.assertRaises(ProfileRenderCancelled):
                save_profile_report_atomic(result, target, check_cancelled=cancel_after_save)
            self.assertFalse(target.exists())
            self.assertEqual(list(Path(temporary).glob("*.tmp.docx")), [])


def _template_for(
    report_type: str,
    *,
    mixed_anchors: bool = False,
) -> tuple[bytes, dict[str, dict[str, str]]]:
    document = Document()
    anchors: dict[str, dict[str, str]] = {}
    for index, semantic in enumerate(REPORT_REQUIREMENTS[report_type]):
        token = "{{" + semantic.upper().replace(".", "_") + "}}"
        if mixed_anchors and index == 0:
            paragraph = document.add_paragraph()
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "1")
            start.set(qn("w:name"), "report_title_anchor")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "1")
            paragraph._p.append(start)
            paragraph.add_run("Title placeholder")
            paragraph._p.append(end)
            anchors[semantic] = {"kind": "bookmark", "value": "report_title_anchor"}
        elif mixed_anchors and index == 1:
            control = OxmlElement("w:sdt")
            properties = OxmlElement("w:sdtPr")
            tag = OxmlElement("w:tag")
            tag.set(qn("w:val"), "overview_anchor")
            properties.append(tag)
            content = OxmlElement("w:sdtContent")
            paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "Overview placeholder"
            run.append(text)
            paragraph.append(run)
            content.append(paragraph)
            control.append(properties)
            control.append(content)
            document.element.body.insert(len(document.element.body) - 1, control)
            anchors[semantic] = {"kind": "content_control", "value": "overview_anchor"}
        else:
            document.add_paragraph(token)
            anchors[semantic] = {"kind": "token", "value": token}
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), anchors


def _published_pack(
    template_bytes: bytes,
    report_type: str,
    anchors: dict[str, dict[str, str]],
) -> bytes:
    with tempfile.TemporaryDirectory() as temporary:
        service = TemplateStudioService(Path(temporary) / "studio")
        workspace = _complete_workspace(service, template_bytes, report_type, anchors)
        return build_template_pack(workspace, template_bytes, _evidence())


def _complete_workspace(
    service: TemplateStudioService,
    template_bytes: bytes,
    report_type: str,
    anchors: dict[str, dict[str, str]],
) -> dict:
    workspace = service.create(
        template_bytes,
        report_type=report_type,
        profile_id=f"renderer-{report_type.replace('_', '-')}",
        display_name=f"Renderer {report_type}",
    )
    for semantic in REPORT_REQUIREMENTS[report_type]:
        requirement = COMMON_REQUIREMENTS[semantic]
        workspace = service.approve(
            workspace["workspaceId"],
            semantic=semantic,
            anchor=anchors[semantic],
            fields=[
                {"source": source, "target": f"column:{index}"}
                for index, source in enumerate(requirement.required_fields, start=1)
            ],
            expected_revision=workspace["revision"],
        )
    return workspace


def _prepared(
    report_type: str,
    template_bytes: bytes,
    payload: dict,
) -> PreparedReportSnapshot:
    accepted = AcceptedReportSnapshot.create(
        rows=[],
        metadata={},
        title="Profile Renderer Test",
        organization="Reporter Pro",
        assessment_date="2026-08-28",
        report_type=report_type,
        template_bytes=template_bytes,
        template_key="profile-renderer-test",
    )
    return PreparedReportSnapshot.create(accepted, payload=payload)


def _evidence() -> TemplatePackValidationEvidence:
    return TemplatePackValidationEvidence(
        fixture_id="profile-renderer-fixture",
        fixture_passed=True,
        integrity_passed=True,
        visual_approved=True,
        validator_version="test-1.0",
        validated_at="2026-08-28T00:00:00Z",
        validation_run_id="1" * 64,
        artifact_sha256="2" * 64,
        structural_sha256="3" * 64,
        reviewed_by="profile-renderer-test",
        reviewed_at="2026-08-28T00:00:00Z",
    )


if __name__ == "__main__":
    unittest.main()
